# SPDX-License-Identifier: AGPL-3.0-only
import hashlib
import importlib.util
import json
import math
import tempfile
import unittest
import zipfile
from pathlib import Path
import io
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/"engine"))
import anasing_engine as E
from PIL import Image
import fitz

class WorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix="Anasing-integration-");cls.root=Path(cls.temp.name)
        cls.inputs=cls.root/"inputs";cls.inputs.mkdir();cls.outputs=cls.root/"outputs";cls.outputs.mkdir()
        cls.im=cls.inputs/"image.png"
        image=Image.new("RGB",(320,240),"#66AAEE");image.putpixel((0,0),(255,0,0));image.save(cls.im)
        cls.other=cls.inputs/"second.png";Image.new("RGB",(180,300),"#EEEE55").save(cls.other)
        cls.svg=cls.inputs/"vector.svg";cls.svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="320" height="240"><rect width="320" height="240" fill="red"/></svg>')
        cls.pdf=cls.inputs/"document.pdf";pdf=fitz.open()
        for i in range(2): pdf.new_page().insert_text((40,70),"Anasing integration page %d\nAmount 42\nHello world" % (i+1))
        pdf.save(cls.pdf);pdf.close()
        cls.audio=cls.inputs/"recording.wav";E.ff(["-f","lavfi","-i","sine=frequency=440:duration=2",str(cls.audio)])
        cls.video=cls.inputs/"movie.mp4";E.ff(["-f","lavfi","-i","testsrc2=size=320x240:rate=15:duration=2","-f","lavfi","-i","sine=frequency=440:duration=2","-c:v","libx264","-pix_fmt","yuv420p","-c:a","aac","-shortest",str(cls.video)])
        cls.docx=cls.inputs/"notes.docx"
        from docx import Document
        doc=Document();doc.add_heading("Anasing document",0);doc.add_paragraph("A useful test of local conversions.");doc.save(cls.docx)
        cls.xlsx=cls.inputs/"table.xlsx"
        from openpyxl import Workbook
        wb=Workbook();wb.active.append(["Item","Amount"]);wb.active.append(["Apples",42]);wb.save(cls.xlsx)
        cls.pptx=cls.inputs/"slides.pptx"
        from pptx import Presentation
        from pptx.util import Inches
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2)).text="Anasing slide";ppt.save(cls.pptx)
        cls.md=cls.inputs/"note.md";cls.md.write_text("# Anasing\n\nHello **world**.\n")
        cls.txt=cls.inputs/"note.txt";cls.txt.write_text("Hello from Anasing.\nSecond line.")
        cls.csv=cls.inputs/"table.csv";cls.csv.write_text('Name,Value\n"A, B",42\nC,17\n')
        cls.tsv=cls.inputs/"table.tsv";cls.tsv.write_text('Name\tValue\nA\t42\n')
        cls.srt=cls.inputs/"captions.srt";cls.srt.write_text("1\n00:00:00,000 --> 00:00:01,000\nHello world.\n\n2\n00:00:01,000 --> 00:00:02,000\nSecond cue.\n")
        cls.archive=cls.inputs/"files.zip"
        with zipfile.ZipFile(cls.archive,"w") as z: z.writestr("folder/hello.txt","Keep the hierarchy")
        cls.hashes={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in cls.inputs.iterdir()}
    @classmethod
    def tearDownClass(cls):
        for p,checksum in cls.hashes.items():
            if hashlib.sha256(p.read_bytes()).hexdigest()!=checksum: raise AssertionError("Original changed: "+str(p))
        cls.temp.cleanup()
    def op(self,action,sources=None,**parameters):
        result=E.execute({"action":action,"sources":[str(s) for s in (sources or [self.im])],"outputDirectory":str(self.outputs),"parameters":parameters})
        self.assertTrue(result["ok"]);paths=[Path(p) for p in result["outputs"]]
        for path in paths:
            self.assertTrue(path.exists())
            if path.is_file(): self.assertGreater(path.stat().st_size,0)
        return paths[0]
    def test_image_conversions(self):
        for target in ["png","jpg","webp","heic","tiff","avif","bmp","pdf"]:
            with self.subTest(target=target):
                out=self.op("convert:"+target)
                if target=="pdf": self.assertEqual(len(fitz.open(out)),1)
                else:
                    loaded=Image.open(out);self.assertEqual(loaded.size,(320,240))
        self.assertEqual(Image.open(self.op("convert:png",[self.svg])).size,(640,480))
    def test_image_tools(self):
        self.assertEqual(Image.open(self.op("cropImage",region=[.25,.25,.5,.5])).size,(160,120))
        redacted=Image.open(self.op("redactImage",region=[.25,.25,.5,.5]));self.assertEqual(redacted.convert("RGB").getpixel((160,120)),(0,0,0))
        edited=Image.open(self.op("editImage",brightness=.5,rotation=90));self.assertEqual(edited.size,(240,320))
        self.assertEqual(Image.open(self.op("frameImage",margin=20)).size,(360,280))
        self.op("annotateImage",text="Hello",region=[.1,.1,.5,.5])
        self.assertEqual(Image.open(self.op("compress",width=160)).size,(160,120))
        tagged=self.op("metadata",metadata={"title":"Anasing title"});self.assertEqual(Image.open(tagged).info.get("title"),"Anasing title")
        clean=self.op("metadata",[tagged]);self.assertNotIn("title",Image.open(clean).info)
        self.assertEqual(len(fitz.open(self.op("createPDF",[self.im,self.other]))),2)
        self.assertEqual(Image.open(self.op("collage",[self.im,self.other],columns=2,cell=100,margin=10)).size,(230,120))
    def test_audio_formats(self):
        for target in ["mp3","m4a","wav","flac","ogg","opus","aiff","wma"]:
            with self.subTest(target=target):
                out=self.op("convert:"+target,[self.audio]);info=E.probe(out)
                self.assertTrue(any(s["codec_type"]=="audio" for s in info["streams"]))
    def test_audio_tools(self):
        for tool,params in [("compress",{}),("metadata",{"metadata":{"title":"Anasing audio"}}),("normalizeAudio",{}),("visualize",{}),("trimAudio",{"start":.5,"end":1.5}),("channels",{"channels":2}),("bleep",{"start":.3,"end":.9})]:
            with self.subTest(tool=tool):
                out=self.op(tool,[self.audio],**params);info=E.probe(out)
                if tool=="trimAudio": self.assertAlmostEqual(float(info["format"]["duration"]),1,delta=.1)
                if tool=="channels": self.assertEqual(info["streams"][0]["channels"],2)
                if tool=="visualize": self.assertTrue(any(s["codec_type"]=="video" for s in info["streams"]))
    def test_video_formats(self):
        for target in ["mp4","mov","mkv","webm","avi","wmv","gif","mp3"]:
            with self.subTest(target=target):
                info=E.probe(self.op("convert:"+target,[self.video]));self.assertGreater(len(info["streams"]),0)
    def test_video_tools(self):
        for tool,params in [("compress",{}),("metadata",{"metadata":{"title":"Anasing video"}}),("muteVideo",{}),("trimVideo",{"start":.5,"end":1.5}),("cropVideo",{"region":[.25,.25,.5,.5]}),("speed",{"speed":2}),("snapshot",{"times":"0,1"}),("splitVideo",{"seconds":1}),("redactVideo",{"region":[.1,.1,.5,.5],"start":0,"end":1})]:
            with self.subTest(tool=tool):
                out=self.op(tool,[self.video],**params)
                if tool=="snapshot": self.assertEqual(len(list(out.glob("*.png"))),2)
                elif tool=="splitVideo": self.assertEqual(len(list(out.glob("*.mp4"))),2)
                else:
                    info=E.probe(out);video=next(s for s in info["streams"] if s["codec_type"]=="video")
                    if tool=="muteVideo": self.assertFalse(any(s["codec_type"]=="audio" for s in info["streams"]))
                    if tool=="cropVideo": self.assertEqual((video["width"],video["height"]),(160,120))
                    if tool=="speed": self.assertAlmostEqual(float(info["format"]["duration"]),1,delta=.2)
        joined=self.op("joinVideos",[self.video,self.video]);self.assertAlmostEqual(float(E.probe(joined)["format"]["duration"]),4,delta=.3)
    def test_pdf_exports(self):
        for target in ["docx","pptx","xlsx","jpg","png","epub","txt","md"]:
            with self.subTest(target=target):
                out=self.op("convert:"+target,[self.pdf])
                if target in {"jpg","png"}: self.assertEqual(len(list(out.iterdir())),2)
                elif target in {"docx","pptx","xlsx","epub"}: self.assertTrue(zipfile.is_zipfile(out))
                else: self.assertIn("Anasing integration",out.read_text())
    def test_pdf_tools(self):
        self.assertEqual(len(list(self.op("splitPDF",[self.pdf]).glob("*.pdf"))),2)
        self.assertEqual(len(fitz.open(self.op("mergePDF",[self.pdf,self.pdf]))),4)
        order=fitz.open(self.op("organizePDF",[self.pdf],pages="2,1",rotation=90));self.assertIn("page 2",order[0].get_text());self.assertEqual(order[0].rotation,90)
        self.assertEqual(fitz.open(self.op("metadata",[self.pdf],metadata={"title":"Test"})).metadata["title"],"Test")
        self.assertEqual(len(fitz.open(self.op("compress",[self.pdf]))),2)
    def test_documents(self):
        for source in [self.docx,self.md,self.txt]:
            for target in ["docx","pdf","txt","md","html","epub","png","jpg"]:
                with self.subTest(source=source.suffix,target=target):
                    out=self.op("convert:"+target,[source])
                    if target=="pdf": self.assertGreater(len(fitz.open(out)),0)
                    if target in {"png","jpg"}: self.assertGreater(len(list(out.iterdir())),0)
        epub=self.op("convert:epub",[self.docx])
        for target in ["docx","pdf","txt","md","html"]:
            with self.subTest(epub=target): self.op("convert:"+target,[epub])
    def test_slides(self):
        for target in ["pdf","jpg","png"]:
            with self.subTest(target=target):
                out=self.op("convert:"+target,[self.pptx]);self.assertTrue(out.exists())
    def test_spreadsheets(self):
        for source in [self.xlsx,self.csv,self.tsv]:
            for target in ["xlsx","ods","csv","tsv","pdf"]:
                with self.subTest(source=source.suffix,target=target):
                    out=self.op("convert:"+target,[source])
                    if target=="xlsx":
                        from openpyxl import load_workbook
                        wb=load_workbook(out);self.assertGreaterEqual(wb.active.max_row,2)
        output=self.op("convert:xlsx",[self.csv])
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(output).active["A2"].value,"A, B")
    def test_subtitles(self):
        vtt=self.op("convert:vtt",[self.srt]);self.assertIn("00:00:01.000",vtt.read_text());self.assertTrue(vtt.read_text().startswith("WEBVTT"))
        srt=self.op("convert:srt",[vtt]);self.assertIn("00:00:01,000",srt.read_text())
        txt=self.op("convert:txt",[vtt]);self.assertIn("Hello world",txt.read_text());self.assertNotIn("-->",txt.read_text())
    def test_archives_and_safety(self):
        extracted=self.op("extract",[self.archive]);self.assertEqual((extracted/"folder/hello.txt").read_text(),"Keep the hierarchy")
        for target in ["zip","tar","gz"]:
            with self.subTest(target=target):
                archive=self.op("convert:"+target,[self.archive]);out=self.op("extract",[archive]);self.assertEqual((out/"folder/hello.txt").read_text(),"Keep the hierarchy")
        evil=self.root/"evil.zip"
        with zipfile.ZipFile(evil,"w") as z: z.writestr("../escape.txt","bad")
        with self.assertRaisesRegex(ValueError,"unsafe path"): self.op("extract",[evil])
        self.assertFalse((self.outputs/"escape.txt").exists())
    def test_scanned_pdf_ocr(self):
        from PIL import ImageDraw,ImageFont
        raster=Image.new("RGB",(1000,400),"white")
        ImageDraw.Draw(raster).text((50,100),"Anasing scanned document",font=ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf",48),fill="black")
        buffer=io.BytesIO();raster.save(buffer,format="PNG")
        source=self.root/"scan.pdf";doc=fitz.open();page=doc.new_page(width=500,height=200);page.insert_image(page.rect,stream=buffer.getvalue());doc.save(source);doc.close()
        output=self.op("convert:txt",[source]);self.assertIn("Anasing scanned document",output.read_text())
    def test_cancel_cleans_temporary_work(self):
        import subprocess,time
        request={"action":"joinVideos","sources":[str(self.video)]*100,"outputDirectory":str(self.outputs),"parameters":{}}
        process=subprocess.Popen([sys.executable,str(Path(E.__file__))],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=subprocess.PIPE)
        process.stdin.write(json.dumps(request).encode());process.stdin.close();process.stdin=None
        time.sleep(.7)
        self.assertIsNone(process.poll())
        process.terminate();process.communicate(timeout=10)
        self.assertFalse(list(self.outputs.glob(".anasing-*")))
        self.assertFalse(list(self.outputs.glob("*joinVideos*")))
    def test_default_output_beside_each_original(self):
        import shutil
        a=(self.root/"source A").resolve();b=(self.root/"source B").resolve();a.mkdir();b.mkdir()
        first=a/"photo.png";second=b/"photo.png"
        shutil.copy2(self.im,first);shutil.copy2(self.other,second)
        original=first.read_bytes()
        result=E.execute({"action":"editImage","sources":[str(first),str(second)],"parameters":{"brightness":.5}})
        paths=[Path(p) for p in result["outputs"]]
        self.assertEqual([p.parent for p in paths],[a,b])
        self.assertTrue(all(p.name=="photo — Edited.png" for p in paths))
        self.assertEqual(first.read_bytes(),original)
        self.assertNotEqual(paths[0].read_bytes(),original)
        conversion=E.execute({"action":"convert:jpg","sources":[str(first)],"parameters":{}})
        self.assertEqual(Path(conversion["outputs"][0]).parent,a)
        combined=E.execute({"action":"createPDF","sources":[str(first),str(second)],"parameters":{}})
        self.assertEqual(Path(combined["outputs"][0]).parent,a)
        self.assertFalse(list(a.glob(".anasing-*")))
    def test_invalid_requests_and_collision(self):
        with self.assertRaisesRegex(ValueError,"Pages export"): self.op("convert:pages",[self.docx])
        with self.assertRaises(ValueError): self.op("cropImage",region=[-1,0,1,1])
        with self.assertRaises(ValueError): self.op("trimAudio",[self.audio],start=10,end=1)
        a=self.op("convert:jpg");old=a.read_bytes();b=self.op("convert:jpg");self.assertNotEqual(a,b);self.assertEqual(a.read_bytes(),old)
        self.assertFalse(list(self.outputs.glob(".anasing-*")))

if __name__=="__main__": unittest.main(verbosity=2)
